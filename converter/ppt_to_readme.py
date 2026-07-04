"""
PPT → README (Markdown) Converter
===================================
Converts PowerPoint (.pptx) presentations into structured Markdown files.

Extracts:
    - Slide titles as ## headings
    - Bullet text as markdown lists
    - Speaker notes as blockquotes
    - Embedded images exported to a subfolder and linked
    - Tables as markdown tables

Entry points:
    run()   — launched from the project's converter/main.py submenu
    main()  — direct CLI usage:  python ppt_to_readme.py /path/to/file.pptx
"""

import os
import sys
import re
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
except ImportError:
    print("[ERROR] python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)

from shared.console import (
    console,
    print_banner,
    print_success,
    print_error,
    print_warning,
    print_info,
)
from shared.config import PROJECT_ROOT

from rich.panel import Panel
from rich.progress import (
    Progress,
    SpinnerColumn,
    BarColumn,
    TextColumn,
    TimeElapsedColumn,
)
from rich.prompt import Prompt


# ---------------------------------------------------------------------------
# Core Conversion Engine
# ---------------------------------------------------------------------------
def _sanitize_filename(name: str) -> str:
    """Create a filesystem-safe filename from a string."""
    name = re.sub(r'[<>:"/\\|?*]', '_', name)
    name = re.sub(r'\s+', '_', name)
    return name.strip('_')[:100] or "untitled"


def _extract_slide_text(shape) -> list[str]:
    """Extract text lines from a shape's text frame."""
    lines = []
    if not shape.has_text_frame:
        return lines

    for para in shape.text_frame.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Detect bullet level
        level = para.level if para.level else 0
        indent = "  " * level

        lines.append(f"{indent}- {text}")

    return lines


def _extract_table_markdown(shape) -> str:
    """Convert a PowerPoint table shape to markdown table syntax."""
    if not shape.has_table:
        return ""

    table = shape.table
    rows = []

    for row in table.rows:
        cells = []
        for cell in row.cells:
            cell_text = cell.text.strip().replace("|", "\\|")
            cells.append(cell_text)
        rows.append(cells)

    if not rows:
        return ""

    # Build markdown table
    md_lines = []
    # Header row
    md_lines.append("| " + " | ".join(rows[0]) + " |")
    # Separator
    md_lines.append("| " + " | ".join(["---"] * len(rows[0])) + " |")
    # Data rows
    for row in rows[1:]:
        # Pad row if needed
        while len(row) < len(rows[0]):
            row.append("")
        md_lines.append("| " + " | ".join(row[:len(rows[0])]) + " |")

    return "\n".join(md_lines)


def _extract_notes(slide) -> str:
    """Extract speaker notes from a slide."""
    if not slide.has_notes_slide:
        return ""

    notes_text = slide.notes_slide.notes_text_frame.text.strip()
    if not notes_text:
        return ""

    # Format as blockquote
    lines = notes_text.split("\n")
    quoted = "\n".join(f"> {line}" for line in lines)
    return quoted


def convert_pptx_to_readme(
    pptx_path: Path,
    output_dir: Path | None = None,
    progress_cb=None,
) -> Path:
    """
    Convert a .pptx file to a structured Markdown README.

    Parameters
    ----------
    pptx_path : Path
        Path to the .pptx file.
    output_dir : Path, optional
        Directory for output. Defaults to a folder named after the PPTX next to it.
    progress_cb : callable, optional
        Called as progress_cb(current_slide, total_slides).

    Returns
    -------
    Path
        The path to the generated README.md file.
    """
    pptx_path = Path(pptx_path).resolve()
    if not pptx_path.exists():
        raise FileNotFoundError(f"File not found: {pptx_path}")
    if pptx_path.suffix.lower() not in (".pptx",):
        raise ValueError(f"Not a PowerPoint file: {pptx_path.name}")

    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)
    total_slides = len(slides)

    if total_slides == 0:
        raise ValueError("Presentation has no slides.")

    # Setup output directory
    pptx_stem = pptx_path.stem
    if output_dir is None:
        output_dir = pptx_path.parent / _sanitize_filename(pptx_stem)

    output_dir.mkdir(parents=True, exist_ok=True)
    images_dir = output_dir / "images"
    images_dir.mkdir(exist_ok=True)

    # Extract presentation title from first slide (if available)
    pres_title = pptx_stem.replace("_", " ").replace("-", " ").title()
    if slides:
        first_slide = slides[0]
        for shape in first_slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                pres_title = shape.text.strip()
                break

    # Build markdown content
    md_sections = []
    md_sections.append(f"# {pres_title}\n")
    md_sections.append(f"*Converted from `{pptx_path.name}` — {total_slides} slides*\n")
    md_sections.append("---\n")

    # Table of contents
    toc_lines = ["## Table of Contents\n"]
    for idx, slide in enumerate(slides, 1):
        slide_title = f"Slide {idx}"
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                # Use the first text shape as slide title
                candidate = shape.text.strip().split("\n")[0][:80]
                if candidate:
                    slide_title = candidate
                    break
        anchor = re.sub(r'[^a-z0-9 -]', '', slide_title.lower())
        anchor = re.sub(r'\s+', '-', anchor)
        toc_lines.append(f"- [Slide {idx}: {slide_title}](#{f'slide-{idx}-{anchor}'})")

    md_sections.append("\n".join(toc_lines))
    md_sections.append("\n---\n")

    # Process each slide
    image_counter = 0
    for idx, slide in enumerate(slides, 1):
        # Find slide title
        slide_title = f"Slide {idx}"
        title_found = False
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip() and not title_found:
                candidate = shape.text.strip().split("\n")[0][:80]
                if candidate:
                    slide_title = candidate
                    title_found = True

        md_sections.append(f"\n## Slide {idx}: {slide_title}\n")

        # Process shapes
        for shape in slide.shapes:
            # Skip the title shape we already used
            if title_found and shape.has_text_frame:
                first_line = shape.text.strip().split("\n")[0][:80]
                if first_line == slide_title:
                    # Extract remaining text (non-title lines)
                    remaining = _extract_slide_text(shape)
                    # Skip the first item (title) if it matches
                    if remaining and remaining[0].lstrip("- ").strip() == slide_title:
                        remaining = remaining[1:]
                    if remaining:
                        md_sections.append("\n".join(remaining) + "\n")
                    continue

            # Images
            if shape.shape_type == 13:  # Picture
                try:
                    image_counter += 1
                    image_blob = shape.image.blob
                    content_type = shape.image.content_type
                    ext = ".png"
                    if "jpeg" in content_type or "jpg" in content_type:
                        ext = ".jpg"
                    elif "gif" in content_type:
                        ext = ".gif"
                    elif "bmp" in content_type:
                        ext = ".bmp"

                    img_filename = f"slide_{idx:02d}_img_{image_counter:03d}{ext}"
                    img_path = images_dir / img_filename
                    with open(img_path, "wb") as f:
                        f.write(image_blob)

                    md_sections.append(f"\n![Slide {idx} Image](images/{img_filename})\n")
                except Exception:
                    md_sections.append("\n*[Image could not be extracted]*\n")

            # Text content
            elif shape.has_text_frame:
                text_lines = _extract_slide_text(shape)
                if text_lines:
                    md_sections.append("\n".join(text_lines) + "\n")

            # Tables
            elif shape.has_table:
                table_md = _extract_table_markdown(shape)
                if table_md:
                    md_sections.append(f"\n{table_md}\n")

        # Speaker notes
        notes = _extract_notes(slide)
        if notes:
            md_sections.append(f"\n**Speaker Notes:**\n\n{notes}\n")

        md_sections.append("\n---\n")

        if progress_cb:
            progress_cb(idx, total_slides)

    # Write README.md
    readme_path = output_dir / "README.md"
    full_md = "\n".join(md_sections)

    with open(readme_path, "w", encoding="utf-8") as f:
        f.write(full_md)

    return readme_path


# ---------------------------------------------------------------------------
# CLI Conversion with Rich progress
# ---------------------------------------------------------------------------
def run_cli_conversion(pptx_path: Path):
    """Run PPT → README conversion in terminal mode with rich progress display."""
    pptx_path = Path(pptx_path).resolve()

    if not pptx_path.exists():
        print_error(f"File not found: {pptx_path}")
        return
    if pptx_path.suffix.lower() != ".pptx":
        print_error(f"Not a PPTX file: {pptx_path.name}")
        return

    console.print()
    console.print(
        Panel(
            f"[bold info]PPT → README Converter[/bold info]\n"
            f"File: [yellow]{pptx_path.name}[/yellow]",
            border_style="cyan",
        )
    )

    import time
    start = time.time()

    with Progress(
        SpinnerColumn(),
        TextColumn("[progress.description]{task.description}"),
        BarColumn(bar_width=40, complete_style="cyan", finished_style="green"),
        TextColumn("[progress.percentage]{task.percentage:>3.0f}%"),
        TimeElapsedColumn(),
        console=console,
    ) as progress:
        task = progress.add_task("[cyan]Extracting slides...", total=100)

        def rich_progress(curr, total):
            pct = (curr / total) * 100 if total else 0
            progress.update(
                task,
                completed=pct,
                description=f"[cyan]Slide {curr}/{total}",
            )

        try:
            result = convert_pptx_to_readme(pptx_path, progress_cb=rich_progress)
            elapsed = time.time() - start

            console.print()
            console.print(
                Panel(
                    f"[success][+][/success] [bold green]Conversion completed![/bold green]\n\n"
                    f"  • [info]Output:[/] {result}\n"
                    f"  • [info]Time Taken:[/] {elapsed:.2f} seconds",
                    border_style="green",
                )
            )
        except KeyboardInterrupt:
            console.print("\n  [warning][*][/warning] Conversion interrupted.")
        except Exception as e:
            console.print(f"\n  [error][!][/error] Failed to convert: {e}")


# ---------------------------------------------------------------------------
# Interactive TUI
# ---------------------------------------------------------------------------
def run_interactive_tui():
    """Terminal UI prompted launcher for PPT → README."""
    console.print(
        Panel(
            "[bold info]PPT → README Converter[/bold info]\n\n"
            "  Converts a PowerPoint presentation to structured Markdown.\n"
            "  Extracts titles, bullets, tables, images, and speaker notes.\n\n"
            "  1. Convert a PPTX file to README\n"
            "  0. Back / Exit",
            title="[bold]PPT → README[/bold]",
            border_style="magenta",
            padding=(1, 3),
        )
    )

    choice = Prompt.ask("  Choice", choices=["0", "1"], default="1")
    if choice == "0":
        return
    elif choice == "1":
        path_str = Prompt.ask("  Enter PPTX file path").strip()
        if path_str:
            p = Path(path_str)
            if p.is_file() and p.suffix.lower() == ".pptx":
                run_cli_conversion(p)
            else:
                print_error(f"Not a valid PPTX file: '{path_str}'")
        else:
            print_warning("Empty path, returning.")


# ---------------------------------------------------------------------------
# Entry points
# ---------------------------------------------------------------------------
def run():
    """Entry point for the project's converter/main.py submenu."""
    run_interactive_tui()


def main():
    """Direct CLI entry point."""
    if len(sys.argv) > 1:
        path_arg = Path(sys.argv[1])
        if path_arg.is_file() and path_arg.suffix.lower() == ".pptx":
            run_cli_conversion(path_arg)
        else:
            print(f"[ERROR] Not a valid PPTX file: {sys.argv[1]}")
    else:
        run_interactive_tui()


if __name__ == "__main__":
    main()
